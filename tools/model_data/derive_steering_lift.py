"""Derive the front steering-lift coefficients from steering-axis hardpoints.

Rotating the upright about the line through the upper and lower outer ball joints
carries the wheel centre along a circle, so its height changes with the steer
angle. This tool computes that height ``h(delta)`` (positive = wheel centre moves
up relative to the body) over a symmetric steer range and fits the two-term model
used by the 14DoF suspension kinematics::

    h(delta) = b * sin(delta) + a * (1 - cos(delta))

``delta`` is the road-wheel angle, positive for a right turn. The right-hand
corner is computed from the hardpoints and the left-hand corner by mirroring
``y``, so ``b`` flips sign between the corners and ``a`` does not.

The hardpoints are read from a workbook given on the command line (real car
data). Rows are found by label, not by cell address: the block whose header
cell reads ``Front Right Hardpoints`` is followed by labelled rows
(``uca outer``, ``lca outer``, ``wheel center``) with X, Y, Z in the next three
columns. The coordinate system is the CAD one: x forward, y to the right, z
DOWN, in inches. Pass ``--units-per-metre`` for other units.

Only the steering-axis rotation of the wheel centre is modelled. Jounce
dependence, compliance and the small camber change (which moves the contact
patch by about ``R(1 - cos(gamma))`` ~ 0.1 mm at full lock) are neglected.

Usage::

    PYTHONPATH=src:. python3 tools/model_data/derive_steering_lift.py \\
        --workbook <hardpoint workbook> --sheet <sheet name> --max-steer-deg 30
"""

from __future__ import annotations

import argparse
import sys
from dataclasses import dataclass

import numpy as np

INCH = 0.0254
FRONT_HEADER = "Front Right Hardpoints"
LABELS = ("uca outer", "lca outer", "wheel center")


class SteeringLiftError(ValueError):
    """Raised when hardpoints are missing or the geometry is degenerate."""


@dataclass(frozen=True)
class SteeringAxis:
    """Right-hand steering geometry in metres: x forward, y right, z DOWN."""

    upper: np.ndarray
    lower: np.ndarray
    wheel_centre: np.ndarray

    def __post_init__(self):
        for name in ("upper", "lower", "wheel_centre"):
            value = np.asarray(getattr(self, name), dtype=float)
            if value.shape != (3,) or not np.all(np.isfinite(value)):
                raise SteeringLiftError(f"{name} must be a finite 3-vector")
            object.__setattr__(self, name, value)
        if np.linalg.norm(self.lower - self.upper) < 1.0e-6:
            raise SteeringLiftError("upper and lower ball joints coincide")

    def mirrored(self) -> "SteeringAxis":
        flip = np.array([1.0, -1.0, 1.0])
        return SteeringAxis(self.upper*flip, self.lower*flip, self.wheel_centre*flip)


def _rotation(axis: np.ndarray, angle: float) -> np.ndarray:
    n = axis/np.linalg.norm(axis)
    k = np.array([[0.0, -n[2], n[1]], [n[2], 0.0, -n[0]], [-n[1], n[0], 0.0]])
    return np.eye(3) + np.sin(angle)*k + (1.0 - np.cos(angle))*(k @ k)


def _yaw_of(axis: np.ndarray, angle: float) -> float:
    heading = _rotation(axis, angle) @ np.array([1.0, 0.0, 0.0])
    return float(np.arctan2(heading[1], heading[0]))


def _axis_angle_for_yaw(axis: np.ndarray, yaw: float) -> float:
    """Return the rotation about ``axis`` that yaws the wheel heading by ``yaw``."""

    if yaw == 0.0:
        return 0.0
    low, high = -1.5, 1.5
    f_low, f_high = _yaw_of(axis, low) - yaw, _yaw_of(axis, high) - yaw
    if f_low*f_high > 0.0:
        raise SteeringLiftError(f"road-wheel yaw {yaw:.3f} rad is unreachable about this axis")
    for _ in range(80):
        mid = 0.5*(low + high)
        f_mid = _yaw_of(axis, mid) - yaw
        if f_low*f_mid <= 0.0:
            high = mid
        else:
            low, f_low = mid, f_mid
    return 0.5*(low + high)


def wheel_centre_rise(geometry: SteeringAxis, yaw: float) -> float:
    """Return the wheel-centre rise [m] (up positive) at road-wheel yaw ``yaw`` [rad].

    The lower ball joint is the rotation centre. ``yaw`` is positive for a right
    turn (clockwise from above), which is a positive rotation about the
    downward-pointing upper-to-lower axis in the x-forward, y-right, z-down frame.
    """

    axis = geometry.lower - geometry.upper
    angle = _axis_angle_for_yaw(axis, float(yaw))
    moved = geometry.lower + _rotation(axis, angle) @ (geometry.wheel_centre - geometry.lower)
    return float(-(moved[2] - geometry.wheel_centre[2]))


@dataclass(frozen=True)
class LiftFit:
    sine: float  # b [m]
    cosine: float  # a [m]
    maximum_residual: float  # [m]
    maximum_steer: float  # [rad]


def fit_lift(geometry: SteeringAxis, max_steer: float, samples: int = 121) -> LiftFit:
    """Least-squares fit ``b sin(delta) + a (1 - cos(delta))`` over +-``max_steer``."""

    if not (np.isfinite(max_steer) and max_steer > 0.0):
        raise SteeringLiftError("max_steer must be positive and finite")
    yaw = np.linspace(-max_steer, max_steer, samples)
    rise = np.array([wheel_centre_rise(geometry, value) for value in yaw])
    basis = np.column_stack((np.sin(yaw), 1.0 - np.cos(yaw)))
    (b, a), *_ = np.linalg.lstsq(basis, rise, rcond=None)
    return LiftFit(float(b), float(a), float(np.max(np.abs(basis @ [b, a] - rise))), float(max_steer))


def read_hardpoint_block(workbook_path: str, sheet: str, header_text: str, labels,
                         units_per_metre: float) -> dict[str, np.ndarray]:
    """Read labelled X, Y, Z hardpoints (metres) below one block header, by label.

    ``header_text`` is the cell text that opens the block (for example
    ``Front Right Hardpoints``). The labels sit in that header's column and the
    coordinates in the next three columns, so two blocks side by side that reuse
    the same labels stay separate.
    """

    try:
        import openpyxl
    except ImportError as error:  # pragma: no cover - environment dependent
        raise SteeringLiftError("reading a workbook requires openpyxl") from error
    book = openpyxl.load_workbook(workbook_path, data_only=True, read_only=True)
    try:
        if sheet not in book.sheetnames:
            raise SteeringLiftError(f"no sheet {sheet!r}; found {book.sheetnames}")
        grid = [list(row) for row in book[sheet].iter_rows(values_only=True)]
    finally:
        book.close()
    header = None
    for r, row in enumerate(grid):
        for c, value in enumerate(row):
            if isinstance(value, str) and value.strip().lower() == header_text.lower():
                header = (r, c)
                break
        if header:
            break
    if header is None:
        raise SteeringLiftError(f"no {header_text!r} block found")
    r0, c0 = header
    wanted = {label.lower() for label in labels}
    found: dict[str, np.ndarray] = {}
    for row in grid[r0 + 1:]:
        if c0 >= len(row):
            continue
        label = row[c0]
        if isinstance(label, str) and label.strip().lower() in wanted:
            numbers = row[c0 + 1:c0 + 4]
            if len(numbers) == 3 and all(isinstance(v, (int, float)) for v in numbers):
                found.setdefault(label.strip().lower(), np.array(numbers, dtype=float)/units_per_metre)
    missing = [name for name in labels if name.lower() not in found]
    if missing:
        raise SteeringLiftError(f"missing hardpoint rows: {missing}")
    return found


def read_front_right(workbook_path: str, sheet: str, units_per_metre: float) -> SteeringAxis:
    """Read the front-right steering geometry from a hardpoint workbook by label."""

    found = read_hardpoint_block(workbook_path, sheet, FRONT_HEADER, LABELS, units_per_metre)
    return SteeringAxis(found["uca outer"], found["lca outer"], found["wheel center"])


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    parser.add_argument("--workbook", required=True)
    parser.add_argument("--sheet", required=True)
    parser.add_argument("--max-steer-deg", type=float, default=30.0)
    parser.add_argument("--units-per-metre", type=float, default=1.0/INCH)
    return parser


def main(argv=None) -> int:
    args = build_parser().parse_args(argv)
    try:
        right = read_front_right(args.workbook, args.sheet, args.units_per_metre)
        fit_right = fit_lift(right, np.radians(args.max_steer_deg))
        fit_left = fit_lift(right.mirrored(), np.radians(args.max_steer_deg))
    except (SteeringLiftError, OSError) as error:
        print(f"error: {error}", file=sys.stderr)
        return 2
    print(f"max fit residual [m]: {max(fit_left.maximum_residual, fit_right.maximum_residual):.3e}")
    print(f"steer range +-{args.max_steer_deg:g} deg")
    print("--steering-lift-front "
          f"{fit_left.sine!r} {fit_left.cosine!r} {fit_right.sine!r} {fit_right.cosine!r}"
          "   # b_left a_left b_right a_right [m]")
    return 0


if __name__ == "__main__":
    sys.exit(main())
