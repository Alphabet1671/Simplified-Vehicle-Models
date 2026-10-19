"""Import measured heave and roll mode damper tables into pinned NPZ models.

Source provenance is intentionally recorded only as "from real car data".
"""
from __future__ import annotations

import hashlib
from pathlib import Path

import numpy as np
from openpyxl import load_workbook

from tools.model_data.generate_damper_lut import npz_bytes, process_damper_table

ROOT = Path(__file__).resolve().parents[2]
PRIVATE_DIR = ROOT / "IEM24 Files"


def read_mode_table(workbook_path: Path, sheet: str):
    """Read alternating compression/rebound columns as signed click-major data."""
    book = load_workbook(workbook_path, data_only=True, read_only=True)
    try:
        rows = list(book[sheet].iter_rows(values_only=True))
    finally:
        book.close()
    clicks = np.asarray([rows[1][1 + 2 * i] for i in range(12)], dtype=float)
    speed = np.asarray([rows[r][0] for r in range(5, 17)], dtype=float)
    velocity = np.concatenate((-speed[::-1], [0.0], speed))
    force = []
    for click in range(12):
        comp = np.asarray([rows[r][1 + 2 * click] for r in range(5, 17)], dtype=float)
        rebound = np.asarray([rows[r][2 + 2 * click] for r in range(5, 17)], dtype=float)
        force.append(np.concatenate((rebound[::-1], [0.0], comp)))
    return velocity, clicks, np.asarray(force)


def make_table(filename: str, sheet: str, correction=()):
    source = PRIVATE_DIR / filename
    velocity, clicks, force = read_mode_table(source, sheet)
    return process_damper_table(
        velocity, clicks, force,
        corrections=correction,
        source_sha256="from real car data",
        strict_click_speed_mm_s=50.0,
    )


def write_table(target: Path, arrays):
    data = npz_bytes(arrays)
    target.write_bytes(data)
    return hashlib.sha256(data).hexdigest()


def main():
    # The raw roll curves contain adjacent equal values at +/-10 and +/-25 mm/s.
    config = (
        ("heave", "MultiMatic Heave Dampers.xlsx", "Heave Damper - VC01", ()),
        ("roll", "MultiMatic Roll Dampers.xlsx", "Roll Damper - VC01", ()),
    )
    for mode, filename, sheet, corrections in config:
        arrays = make_table(filename, sheet, corrections)
        target = ROOT / "models/data" / f"iem24_{mode}_damper_vc01.npz"
        digest = write_table(target, arrays)
        print(f"{mode}: {target.relative_to(ROOT)} sha256={digest}")


if __name__ == "__main__":
    main()
